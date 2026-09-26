# -*- coding: utf-8 -*-
"""
`대시보드데이터` · `변동사유(자산)` · `포트폴리오상세` 탭을 수식째로 만들어 xlsx로 떨어뜨린다.
구글 시트에서 [파일 → 가져오기 → 업로드 → 새 시트 삽입]으로 넣으면 된다.

  NAME_MAP="실명1=꿀꿀,실명2=말벌,실명3=째니" python tools/gen_dashboard_tab.py

NAME_MAP이 없으면 이름 치환 없이 만든다(실명이 그대로 CSV에 나가므로 권장하지 않음).
배치·매핑의 근거는 docs/sheet-spec.md. 그 문서와 이 스크립트는 항상 같이 고친다.
"""
import os, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL

SUM_, AST, BUD, ROAD = "'요약(재무제표)'", "'자산관리'", "'예결산관리'", "'로드맵'"
POR, RSN = "'포트폴리오상세'", "'변동사유(자산)'"

PAIRS = [p.split('=', 1) for p in os.environ.get('NAME_MAP', '').split(',') if '=' in p]
if not PAIRS:
    print('! NAME_MAP 미설정 — 이름 치환 없이 생성합니다.', file=sys.stderr)

def alias(expr):
    out = expr
    for real, fake in PAIRS:
        out = 'SUBSTITUTE(%s,"%s","%s")' % (out, real.strip(), fake.strip())
    return out

def txt(ref):
    return '=' + alias('IF(%s="","",%s)' % (ref, ref))

wb = openpyxl.Workbook()
ws = wb.active; ws.title = "대시보드데이터"

HDR, BLK, LBL = PatternFill("solid", fgColor="1E4E8C"), PatternFill("solid", fgColor="0B2E5C"), PatternFill("solid", fgColor="EAF1FA")
W, thin = Font(color="FFFFFF", bold=True, size=10), Side(style="thin", color="C8D4E4")
BD = Border(left=thin, right=thin, top=thin, bottom=thin)

r = 1
def block(t, note=""):
    global r
    c = ws.cell(row=r, column=1, value=t); c.fill = BLK; c.font = W
    if note: ws.cell(row=r, column=2, value=note).font = Font(size=9, color="7A8CA0")
    r += 1
def header(cols):
    global r
    for i, h in enumerate(cols, 1):
        c = ws.cell(row=r, column=i, value=h); c.fill = HDR; c.font = W
        c.alignment = Alignment(horizontal="center"); c.border = BD
    r += 1
def row(vals, numfrom=None, fmt='#,##0'):
    global r
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=i, value=v); c.border = BD
        if numfrom and i > numfrom: c.number_format = fmt
        elif i == 1: c.fill = LBL
    r += 1
def gap():
    global r; r += 1

MON  = [CL(i) for i in range(6, 18)]          # F..Q = 1~12월 (자산관리·예결산관리 공통)
MONH = ["%d월" % m for m in range(1, 13)]

# ── #META ───────────────────────────────────────────────
block("#META", "대시보드 기본 정보 (자동)")
header(["키", "값"])
row(["연도",      "=IFERROR(VALUE(LEFT(%s!B1,4)),2026)" % AST])
row(["기준월번호", '=MAX(1,COUNTIF(%s!F38:Q38,">0"))' % AST])
row(["기준월",    '=B4&"월"'])
row(["제목",      '=B3&"년도 꿀벌이네 자산가계부"'])
row(["갱신일시",  '=TEXT(NOW(),"yyyy-mm-dd HH:mm")'])
YEAR, BASEM = "$B$3", "$B$4"
gap()

# ── #자산 ───────────────────────────────────────────────
ASSETS = [
    ("입출금","자산","현금성","꿀꿀",19), ("국내주식","자산","투자자산","꿀꿀",20),
    ("해외주식","자산","투자자산","꿀꿀",21), ("대체투자","자산","투자자산","꿀꿀",22),
    ("입출금","자산","현금성","말벌",24), ("국내주식","자산","투자자산","말벌",25),
    ("해외주식","자산","투자자산","말벌",26), ("대체투자","자산","투자자산","말벌",27),
    ("연금저축펀드","자산","연금자산","꿀꿀",29), ("연금저축펀드","자산","연금자산","말벌",30),
    ("IRP","자산","연금자산","꿀꿀",31), ("IRP","자산","연금자산","말벌",32),
    ("아파트","자산","부동산","공동",34), ("청약통장","자산","청약저축","꿀꿀",35),
    ("청약통장","자산","청약저축","말벌",36), ("주택담보대출","부채","대출","공동",39),
    ("신용대출","부채","대출","공동",40), ("신용카드대금","부채","대출","공동",41),
]
block("#자산", "자산관리 탭 자동 연동 · 단위: 원")
header(["항목","구분","세부분류","소유자","전년말"] + MONH)
for n, g, s, o, src in ASSETS:
    row([n, g, s, o, "=%s!E%d" % (AST, src)] + ["=%s!%s%d" % (AST, c, src) for c in MON], numfrom=4)
gap()

# ── #손익 ───────────────────────────────────────────────
PL = ([("수익", i) for i in range(15, 20)] + [("변동소비", i) for i in range(25, 37)]
    + [("고정소비", i) for i in range(38, 47)] + [("금융지출", i) for i in range(48, 51)])
block("#손익", "예결산관리 탭 자동 연동 · 단위: 원")
header(["항목","구분","월예산"] + MONH)
for g, src in PL:
    row([txt("%s!E%d" % (BUD, src)), g, "=%s!D%d" % (BUD, src)]
        + ["=IFERROR(%s!%s%d,0)" % (BUD, c, src) for c in MON], numfrom=2)
gap()

# ── #사유 ───────────────────────────────────────────────
RSN_PL = ([(8+i, 15+i) for i in range(5)] + [(16+i, 25+i) for i in range(12)]
        + [(30+i, 38+i) for i in range(9)])
SUMMON = [CL(i) for i in range(16, 28)]       # 요약 P..AA = 1~12월
block("#사유", "요약 탭 우측 메모 + 변동사유(자산) 탭 자동 연동")
header(["항목","구분"] + MONH)
for s_row, b_row in RSN_PL:
    row([txt("%s!E%d" % (BUD, b_row)), "손익"] + [txt("%s!%s%d" % (SUM_, c, s_row)) for c in SUMMON])
for i, (n, g, s, o, src) in enumerate(ASSETS):
    row(["%s (%s)" % (n, o), "자산" if g == "자산" else "부채"]
        + [txt("%s!%s%d" % (RSN, CL(2 + m), 4 + i)) for m in range(12)])
gap()

# ── #월별지표 ───────────────────────────────────────────
block("#월별지표", "월별 핵심 지표 (자동) · 단위: 원")
header(["월","총자산","총부채","순자산","연금자산","순자산(연금제외)","투자자산","현금성자산",
        "수입","지출","저축률","로드맵달성률"])
GOAL = "%s!$F$38*1000" % ROAD                 # 올해(1년차) 목표 기말 금융자산
def metric(label, col, prev=False):
    a, b = "%s!%s" % (AST, col), "%s!%s" % (BUD, col)
    inv = "({a}20+{a}21+{a}22+{a}25+{a}26+{a}27)".format(a=a)
    row([label, "=%s38" % a, "=%s42" % a, "=%s18" % a, "=%s33" % a, "=%s18-%s33" % (a, a),
         "=" + inv, "=%s19+%s24" % (a, a),
         "" if prev else "=IFERROR(%s20,0)" % b,
         "" if prev else "=IFERROR(%s37+%s47+%s51,0)" % (b, b, b),
         "" if prev else '=IFERROR(%s3,"")' % b,
         '=IFERROR(%s/(%s),"")' % (inv, GOAL)], numfrom=1)
metric("전년말", "E", prev=True)
for i, c in enumerate(MON):
    metric("%d월" % (i + 1), c)
for rr in range(r - 13, r):
    ws.cell(row=rr, column=11).number_format = '0.00%'
    ws.cell(row=rr, column=12).number_format = '0.00%'
MFIRST = r - 12                                # 1월 행
gap()

# ── #로드맵 ─────────────────────────────────────────────
block("#로드맵", "로드맵 탭 목표 + 실적 (천원→원 환산)")
header(["연차","연도","목표_기초투자자산","목표_연간증액","목표_전세인상","목표_투자수익",
        "목표_기말금융자산","목표_전세금","목표_순자산","실적_금융자산","달성률","비고"])
for i in range(10):
    src, act, yr = 38 + i, 51 + i, 2026 + i
    live = "INDEX(G%d:G%d,%s)" % (MFIRST, MFIRST + 11, BASEM)
    row([txt("%s!A%d" % (ROAD, src)), yr]
        + ["=%s!%s%d*1000" % (ROAD, c, src) for c in "BCDEFGH"]
        + ["=IF(%d=%s,%s,%s!F%d*1000)" % (yr, YEAR, live, ROAD, act),
           '=IFERROR(J%d/G%d,"")' % (r, r), txt("%s!I%d" % (ROAD, src))], numfrom=2)
for rr in range(r - 10, r):
    ws.cell(row=rr, column=11).number_format = '0.00%'
gap()

# ── #연말스냅샷 ─────────────────────────────────────────
block("#연말스냅샷", "연도별 연말 기준 (2024·2025는 자산관리 탭 보관값)")
header(["연도","총자산","총부채","순자산"])
row([2024, "=%s!U38" % AST, "=%s!U42" % AST, "=%s!U18" % AST], numfrom=1)
row([2025, "=%s!E38" % AST, "=%s!E42" % AST, "=%s!E18" % AST], numfrom=1)
row(['=%s&" (진행중)"' % YEAR] + ["=INDEX(%s%d:%s%d,%s)" % (c, MFIRST, c, MFIRST + 11, BASEM)
                                  for c in "BCD"], numfrom=1)
gap()

# ── #포트폴리오 ─────────────────────────────────────────
block("#포트폴리오", "포트폴리오상세 탭 자동 연동")
header(["구분","세부구분","소유자","종목명","통화","매수원금","평가금액","평가손익","수익률","목표비중"])
for i in range(30):
    src = 4 + i
    row([txt("%s!%s%d" % (POR, c, src)) for c in "ABCFG"]
        + ['=IFERROR(%s!%s%d,"")' % (POR, c, src) for c in "LMNO"]
        + ['=IF(A%d="","",IF(LEFT(A%d,2)="핵심",0.6,IF(LEFT(A%d,2)="위성",0.3,0.1)))' % (r, r, r)],
        numfrom=5)
for rr in range(r - 30, r):
    ws.cell(row=rr, column=9).number_format = '0.00%'
    ws.cell(row=rr, column=10).number_format = '0%'
gap()

# ── #목표텍스트 ─────────────────────────────────────────
block("#목표텍스트", "로드맵 탭 상단 메모")
header(["구분","내용"])
row(["최종목표", txt("%s!G2" % ROAD)])
row(["포트폴리오원칙", txt("%s!J2" % ROAD)])

for i, w in enumerate([22,16,14,12] + [14]*14, 1):
    ws.column_dimensions[CL(i)].width = w
ws.freeze_panes = "A2"

# ══ 변동사유(자산) ═════════════════════════════════════
w2 = wb.create_sheet("변동사유(자산)")
w2["A1"] = "자산·부채 항목 변동 사유"; w2["A1"].font = Font(bold=True, size=12, color="1E4E8C")
w2["A2"] = "※ 특이사항이 있는 달에만 적으시면 됩니다. (수익·비용 사유는 기존처럼 요약 탭에)"
w2["A2"].font = Font(size=9, color="C0392B")
for i, h in enumerate(["항목"] + MONH, 1):
    c = w2.cell(row=3, column=i, value=h); c.fill = HDR; c.font = W
    c.alignment = Alignment(horizontal="center"); c.border = BD
for i, (n, g, s, o, src) in enumerate(ASSETS):
    c = w2.cell(row=4 + i, column=1, value="%s (%s)" % (n, o))
    c.fill = LBL; c.font = Font(bold=True, size=10); c.border = BD
    for j in range(2, 14): w2.cell(row=4 + i, column=j).border = BD
w2.column_dimensions['A'].width = 22
for i in range(2, 14): w2.column_dimensions[CL(i)].width = 26
w2.freeze_panes = "B4"

# ══ 포트폴리오상세 ═════════════════════════════════════
w3 = wb.create_sheet("포트폴리오상세")
w3["A1"] = "환율(USD/KRW)"; w3["A1"].font = Font(bold=True, size=10)
w3["B1"] = '=GOOGLEFINANCE("CURRENCY:USDKRW")'; w3["B1"].number_format = '#,##0.00'
w3["C1"] = "마지막 갱신"; w3["C1"].font = Font(bold=True, size=10)
w3["D1"] = '=TEXT(NOW(),"yyyy-mm-dd HH:mm")'
w3["F1"] = "※ 노란 칸(수량·평단가)만 입력하세요. 현재가·환율은 자동 조회됩니다."
w3["F1"].font = Font(size=9, color="C0392B")
PH = ["구분","세부구분","소유자","계좌","티커(자동조회용)","종목명","통화","수량","평단가(현지통화)",
      "현재가(자동)","적용환율(자동)","매수원금(원)","평가금액(원)","평가손익(원)","수익률","비고"]
for i, h in enumerate(PH, 1):
    c = w3.cell(row=3, column=i, value=h); c.fill = HDR; c.font = W
    c.alignment = Alignment(horizontal="center", wrap_text=True); c.border = BD
YEL, AUT = PatternFill("solid", fgColor="FFF6D6"), PatternFill("solid", fgColor="EAF1FA")
for i in range(30):
    rr = 4 + i
    w3.cell(row=rr, column=10, value='=IFERROR(GOOGLEFINANCE(E%d),"")' % rr)
    w3.cell(row=rr, column=11, value='=IF(G%d="USD",$B$1,1)' % rr)
    w3.cell(row=rr, column=12, value='=IFERROR(H%d*I%d*K%d,"")' % (rr, rr, rr))
    w3.cell(row=rr, column=13, value='=IFERROR(H%d*J%d*K%d,"")' % (rr, rr, rr))
    w3.cell(row=rr, column=14, value='=IFERROR(M%d-L%d,"")' % (rr, rr))
    w3.cell(row=rr, column=15, value='=IFERROR(N%d/L%d,"")' % (rr, rr))
    for j in range(1, 17):
        c = w3.cell(row=rr, column=j); c.border = BD
        if j in (8, 9): c.fill = YEL
        if j in range(10, 16): c.fill = AUT
        if j in range(8, 15): c.number_format = '#,##0.00' if j in (9, 10, 11) else '#,##0'
        if j == 15: c.number_format = '0.00%'
tot = 34
w3.cell(row=tot, column=1, value="합계").font = Font(bold=True, size=10)
for col in (12, 13, 14):
    c = w3.cell(row=tot, column=col, value='=SUM(%s4:%s33)' % (CL(col), CL(col)))
    c.font = Font(bold=True, size=10); c.number_format = '#,##0'; c.border = BD
c = w3.cell(row=tot, column=15, value='=IFERROR(N%d/L%d,"")' % (tot, tot))
c.font = Font(bold=True, size=10); c.number_format = '0.00%'; c.border = BD
for i, w in enumerate([14,14,8,12,20,26,7,10,14,14,12,14,14,14,10,20], 1):
    w3.column_dimensions[CL(i)].width = w
w3.row_dimensions[3].height = 32
w3.freeze_panes = "A4"

out = os.path.join(os.path.dirname(__file__), "대시보드_추가탭.xlsx")
wb.save(out)
print("생성 완료:", out, "| 대시보드데이터", ws.max_row, "행")
